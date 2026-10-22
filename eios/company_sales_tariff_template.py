"""Editable tariff workbook from the selected company's catalogue; no inferred prices."""
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment

from eios.company_tables import MAX_ROWS


def build_template(state):
    articles=state['articles']
    if not articles or len(articles)>MAX_ROWS:
        raise ValueError('La plantilla requiere entre 1 y 5.000 artículos.')
    seen=set()
    for article in articles:
        sku,unit=article.get('article_id'),article.get('unit_as_documented')
        if any(not isinstance(v,str) or not v.strip() for v in (sku,unit)) or sku in seen:
            raise ValueError('Revisa el catálogo: cada SKU debe ser único y tener una unidad documentada.')
        seen.add(sku)
    company=state['info']['company'];cutoff=state['info']['as_of_date']
    workbook=Workbook();sheet=workbook.active;sheet.title='Tarifa'
    headers=['SKU','Descripción (consulta)','Unidad','PVP sin IVA','Moneda','Válida desde','Válida hasta']
    sheet.append(headers)
    for article in articles:
        # Explicit string cells preserve leading zeroes and formula-like source text.
        sheet.append([article['article_id'],article.get('description') or '',article['unit_as_documented'],None,'EUR',None,None])
        for column in (1,2,3,5):
            cell=sheet.cell(sheet.max_row,column);cell.data_type='s';cell.number_format='@'
        sheet.cell(sheet.max_row,4).number_format='#,##0.00'
        for column in (6,7):sheet.cell(sheet.max_row,column).number_format='DD/MM/YYYY'
        for column in (4,6,7):sheet.cell(sheet.max_row,column).fill=PatternFill('solid',fgColor='FFF2CC')
    sheet.freeze_panes='D2';sheet.auto_filter.ref=sheet.dimensions
    for letter,width in zip('ABCDEFG',(24,58,20,20,12,20,20)):sheet.column_dimensions[letter].width=width
    for cell in sheet[1]:
        cell.fill=PatternFill('solid',fgColor='0A369D');cell.font=Font(color='FFFFFF',bold=True)
    sheet['D1'].comment=Comment('PVP recomendado por la unidad indicada en C, sin IVA. Completa con la tarifa de la empresa; no es un precio calculado.','EIOS')
    guide=workbook.create_sheet('Instrucciones')
    instructions=[('EMPRESA',company.get('legal_name',company['company_id'])),('CÓDIGO DE EMPRESA',company['company_id']),
        ('CORTE DEL EXPEDIENTE',cutoff),('NATURALEZA',state['documentary_context']['material_nature']),
        ('1. COMPLETAR','En Tarifa completa las celdas amarillas: PVP sin IVA y fechas de vigencia. Los campos vacíos no equivalen a cero.'),
        ('2. REVISAR','Conserva SKU y unidad como texto. EUR es la moneda propuesta para esta plantilla; confirma la moneda de tu tarifa. Elimina las filas que no vayas a cumplimentar.'),
        ('3. CARGAR','En Carga de archivos usa exactamente esta empresa, naturaleza y corte. Selecciona la hoja Tarifa y el contenido Tarifa de venta: PVP por unidad sin IVA.'),
        ('4. RELACIONAR','SKU → SKU; Unidad → Unidad; PVP recomendado → PVP sin IVA; Moneda → Moneda; Tarifa válida desde → Válida desde; Tarifa válida hasta → Válida hasta. Descripción es solo consulta.'),
        ('5. SELECCIONAR','Comprueba y guarda el archivo. En Propuestas de venta vuelve a seleccionar el artículo y elige expresamente su fila de tarifa.'),
        ('ALCANCE','Plantilla sin precios ni vigencias asignados. No constituye una tarifa aprobada, evidencia admitida ni aceptación de una venta.')]
    for label,value in instructions:
        guide.append([label,str(value)])
        for cell in guide[guide.max_row]:cell.data_type='s';cell.alignment=Alignment(vertical='top',wrap_text=True)
        guide.cell(guide.max_row,1).font=Font(bold=True,color='0A369D')
        guide.row_dimensions[guide.max_row].height=48
    guide.column_dimensions['A'].width=28;guide.column_dimensions['B'].width=110
    output=BytesIO();workbook.save(output);workbook.close()
    safe=re.sub(r'[^A-Za-z0-9_.-]','_',company['company_id'])[:80] or 'empresa'
    return dict(filename='Plantilla_tarifa_'+safe+'.xlsx',content=output.getvalue(),article_count=len(articles))
