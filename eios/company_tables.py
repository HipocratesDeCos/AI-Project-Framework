"""Read ordinary company tables without promoting them to operational evidence."""
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import re
import zipfile

MAX_ROWS = 5000
MAX_COLUMNS = 80
MAX_BYTES = 10 * 1024 * 1024

# Declared company parameters; reception never activates purchasing rules.
STOCK_PARAMETERS = [
    ('minimum_stock','Stock mínimo',False),
    ('safety_stock','Stock de seguridad',False),
    ('reorder_point','Punto de pedido',False),
    ('lead_time_days','Plazo de entrega (días enteros)',False),
]
STOCK_PARAMETER_KEYS = {field for field,_,_ in STOCK_PARAMETERS}

# Fields, display labels, and required fields for assisted reception, not core input.
FIELDS = {
    'supplier': [('supplier_id','Código de proveedor',True),('legal_name','Razón social',True),('tax_id','NIF / identificación fiscal',False),('locality','Localidad',False),('contact_name','Persona de contacto',False),('delivery_location','Lugar de entrega',False)],
    'customer': [('customer_id','Código de cliente',True),('legal_name','Razón social',True),('tax_id','NIF / identificación fiscal',False),('locality','Localidad',False)],
    'article': [('article_id','SKU / código de artículo',True),('description','Descripción',True),('unit','Unidad',True),('supplier_id','Código de proveedor',False),('family','Familia',False)],
    'inventory': [('article_id','SKU / código de artículo',True),('unit','Unidad',True),('physical','Existencias físicas',True),('reserved','Cantidad reservada',False),('observation_date','Fecha del inventario (si figura)',False)]+STOCK_PARAMETERS,
    'stock_policy': [('article_id','SKU / código de artículo',True),('unit','Unidad',True),('observation_date','Fecha de referencia (si figura)',False)]+STOCK_PARAMETERS,
    'purchase_line': [('invoice_number','Número de factura',True),('issue_date','Fecha de factura',True),('supplier_id','Código de proveedor',True),('article_id','SKU / código de artículo',True),('quantity','Cantidad',True),('unit','Unidad',True),('net_amount','Importe neto de la línea',True),('currency','Moneda',True)],
    'sale_price': [('article_id','SKU / código de artículo',True),('unit','Unidad',True),('recommended_sale_price','PVP recomendado por unidad, sin IVA (€)',True),('currency','Moneda',True),('valid_from','Tarifa válida desde',True),('valid_until','Tarifa válida hasta',True)],
    'sales_line': [('invoice_number','Número de factura',True),('issue_date','Fecha de factura',True),('customer_id','Código de cliente',True),('article_id','SKU / código de artículo',True),('quantity','Cantidad',True),('unit','Unidad',True),('net_amount','Importe neto de la línea',True),('currency','Moneda',True)],
    'cash_flow': [('flow_id','Referencia del cobro o pago',True),('direction','COBRO o PAGO',True),('amount','Importe',True),('currency','Moneda',True),('due_date','Vencimiento',True),('counterparty','Cliente / proveedor / concepto',False)],
    'balance': [('account_code','Código de cuenta',True),('description','Descripción',False),('balance','Saldo con signo',True),('currency','Moneda',True),('observation_date','Fecha del saldo (si figura)',False)],
    'document': [],
}
SALES_DETAILS = [
    ('description','Descripción del artículo en factura',False),
    ('line_number','Número de línea',False),
    ('unit_price_before_discount','Precio unitario antes de descuento, sin IVA',False),
    ('discount_percent','Descuento de línea (%)',False),
    ('discount_amount','Descuento total de la línea (€)',False),
    ('vat_percent','Tipo de IVA de la línea (%)',False),
    ('vat_amount','Importe de IVA documentado de la línea',False),
    ('invoice_total','Total de factura documentado, con IVA (mismo valor si se repite)',False),
]
COLLECTION_LINKS = [('customer_id','Código de cliente vinculado',False),('invoice_number','Factura vinculada',False),('installment_ref','Referencia de cuota vinculada',False),('payment_method','Forma de cobro declarada',False)]
FIELDS['sales_line'] += SALES_DETAILS
FIELDS['sales_installment'] = [
    ('installment_ref','Referencia única de cuota',True),('customer_id','Código de cliente',True),
    ('invoice_number','Número de factura',True),('installment_number','Número de cuota',True),
    ('due_date','Fecha de vencimiento',True),('amount','Importe original de la cuota, con IVA',True),
    ('currency','Moneda',True),('payment_method','Forma de cobro',True),
    ('pending_amount','Importe pendiente de cobro al corte (0 si está cobrada)',False),
    ('flow_id','Referencia del cobro ya previsto, si existe',False),
]
FIELDS['cash_flow'] += COLLECTION_LINKS
NUMBERS = {'recommended_sale_price','physical','reserved','quantity','net_amount','amount','balance','line_number','unit_price_before_discount','discount_percent','discount_amount','vat_percent','vat_amount','invoice_total','installment_number','pending_amount'} | STOCK_PARAMETER_KEYS
IDENTIFIERS = {'supplier_id','customer_id','article_id','invoice_number','flow_id','account_code','tax_id','installment_ref'}


from eios.company_workbook_fields import EXTENSIONS, REFERENCE_FIELDS, CACHED_PRICE_FIELDS
for _kind,_fields in EXTENSIONS.items():
    FIELDS[_kind] += [(key,label,False) for key,label in _fields]
FIELDS['price_reference']=REFERENCE_FIELDS
NUMBERS |= {'unit_price','reference_cost','reference_price','target_margin'}
IDENTIFIERS |= {'postal_code','support_ref','order_ref','delivery_ref'}



def tariff_mapping(headers):
    """Suggest only a complete, unique mapping of explicit tariff headers."""
    import unicodedata
    def clean(value):
        text=unicodedata.normalize('NFD',str(value))
        return re.sub('[^a-z0-9]','', ''.join(ch for ch in text if not unicodedata.combining(ch)).lower())
    aliases={
        'article_id':['SKU'], 'unit':['Unidad'], 'recommended_sale_price':['PVP sin IVA'],
        'currency':['Moneda'], 'valid_from':['Válida desde'], 'valid_until':['Válida hasta'],
    }
    result={}
    for field,label,_ in FIELDS['sale_price']:
        names={clean(name) for name in [field,label,*aliases[field]]}
        matches=[i for i,header in enumerate(headers) if clean(header) in names]
        if len(matches)!=1:return None
        result[field]=matches[0]
    return result if len(set(result.values()))==len(result) else None


def _cell(value):
    if isinstance(value,(datetime,date)):
        return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    return value


def read_table(content, filename, sheet=None, header_row=1, encoding='utf-8-sig', delimiter='auto', max_rows=None):
    if max_rows is None:max_rows=MAX_ROWS
    if type(max_rows) is not int or not 1<=max_rows<=20000:
        raise ValueError('Límite de filas no válido.')
    if not content or len(content)>MAX_BYTES:
        raise ValueError('Selecciona un archivo de hasta 10 MiB.')
    if not isinstance(header_row,int) or not 1<=header_row<=100:
        raise ValueError('La fila de cabecera debe estar entre 1 y 100.')
    suffix=filename.lower().rsplit('.',1)[-1]
    if suffix=='pdf':
        if not content.startswith(b'%PDF-'):
            raise ValueError('El archivo no tiene una cabecera PDF válida.')
        return dict(format='pdf',sheets=[],sheet=None,headers=[],rows=[],header_row=header_row)
    if suffix=='csv':
        if encoding not in ('utf-8-sig','cp1252','utf-16'):
            raise ValueError('Codificación no admitida.')
        try: text=content.decode(encoding)
        except UnicodeError: raise ValueError('No se puede leer el CSV con esa codificación. Revisa UTF-8, Windows-1252 o UTF-16.') from None
        if delimiter=='auto':
            try: delimiter=csv.Sniffer().sniff(text[:8192],delimiters=';,\t|').delimiter
            except csv.Error: raise ValueError('No se detecta el separador. Selecciona punto y coma, coma o tabulador.') from None
        if delimiter not in (';',',','\t','|'):
            raise ValueError('Separador CSV no admitido.')
        rows=[]
        try:
            for row in csv.reader(StringIO(text),delimiter=delimiter,strict=True):
                if len(row)>MAX_COLUMNS or len(rows)>=max_rows+header_row:
                    raise ValueError(f'El límite es de {max_rows} filas de datos y 80 columnas.')
                rows.append(row)
        except csv.Error: raise ValueError('El CSV tiene comillas o columnas que no pueden leerse.') from None
        sheet_name='CSV'; sheets=['CSV']
    elif suffix=='xlsx':
        try:
            with zipfile.ZipFile(BytesIO(content)) as z:
                members=z.infolist()
                if len(members)>512 or sum(m.file_size for m in members)>100*1024*1024:
                    raise ValueError('El Excel supera el límite de tamaño descomprimido.')
                for member in members:
                    if member.flag_bits&1 or 'vbaproject' in member.filename.lower():
                        raise ValueError('No se admiten libros cifrados o con macros.')
                    if member.filename.lower().endswith('.xml'):
                        raw=z.read(member)
                        if b'\x00' in raw or b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                            raise ValueError('El XML del Excel contiene una codificación o declaración no admitida.')
            from openpyxl import load_workbook
            workbook=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
            try:
                sheets=workbook.sheetnames
                if not sheets or len(sheets)>100:
                    raise ValueError('El libro debe tener entre 1 y 100 hojas.')
                cached_workbook=load_workbook(BytesIO(content),read_only=True,data_only=True,keep_links=False)
                sheet_name=sheet or sheets[0]
                if sheet_name not in sheets:
                    raise ValueError('La hoja seleccionada no existe.')
                rows=[]
                for row,cached_row in zip(workbook[sheet_name].iter_rows(),cached_workbook[sheet_name].iter_rows()):
                    if len(row)>MAX_COLUMNS or len(rows)>=max_rows+header_row:
                        raise ValueError(f'El límite es de {max_rows} filas de datos y 80 columnas. Exporta una tabla más pequeña.')
                    rows.append([dict(formula=True,cached=_cell(cached_row[index].value)) if c.data_type=='f' else dict(formula=True) if c.data_type=='e' else _cell(c.value) for index,c in enumerate(row)])
            finally:
                workbook.close()
                if 'cached_workbook' in locals():cached_workbook.close()
        except (zipfile.BadZipFile,KeyError):
            raise ValueError('No se puede leer el archivo XLSX.') from None
    else:
        raise ValueError('Admite CSV, XLSX y PDF como soporte. Convierte los libros XLS a XLSX.')
    if len(rows)<header_row:
        raise ValueError('No se encuentra la fila de cabecera.')
    header=rows[header_row-1]
    width=max([len(header)]+[len(r) for r in rows[header_row:]])
    headers=[str(header[i]).strip() if i<len(header) and header[i] is not None else '' for i in range(width)]
    data=[dict(row=i+1,values=row+[None]*(width-len(row))) for i,row in enumerate(rows) if i>=header_row and any(v is not None and str(v).strip() for v in row)]
    if not data:
        raise ValueError('La hoja o CSV no contiene filas de datos.')
    return dict(format=suffix,sheets=sheets,sheet=sheet_name,header_row=header_row,headers=headers,rows=data)


def number(value, convention):
    if isinstance(value,bool) or isinstance(value,dict):
        raise ValueError('Se necesita un número fijo; no una fórmula.')
    if isinstance(value,(int,float,Decimal)):
        result=Decimal(str(value))
    else:
        text=str(value).strip()
        if convention=='ES':
            pattern=r'[+-]?(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?'
            normalized=text.replace('.','').replace(',','.')
        elif convention=='EN':
            pattern=r'[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?'
            normalized=text.replace(',','')
        else: raise ValueError('Selecciona el formato numérico español o internacional.')
        if not re.fullmatch(pattern,text):
            raise ValueError('Número incompatible con el formato seleccionado: '+text[:60])
        result=Decimal(normalized)
    if not result.is_finite() or abs(result)>Decimal('1000000000000000'):
        raise ValueError('Número no finito o fuera de rango.')
    return result


def source_date(value):
    text=str(value).strip()
    try:
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}',text): return date.fromisoformat(text).isoformat()
        if re.fullmatch(r'\d{2}/\d{2}/\d{4}',text): return datetime.strptime(text,'%d/%m/%Y').date().isoformat()
    except ValueError: pass
    raise ValueError('Fecha inválida: usa AAAA-MM-DD o DD/MM/AAAA con año de cuatro cifras.')


def normalize(table, kind, mapping, cutoff, convention):
    if kind not in FIELDS or kind=='document': raise ValueError('Selecciona el tipo de tabla.')
    fields=FIELDS[kind]
    if not isinstance(mapping,dict) or any(key not in {f[0] for f in fields} for key in mapping):
        raise ValueError('Mapeo de columnas no válido.')
    columns=[]
    for field,label,required in fields:
        column=mapping.get(field)
        if column is None:
            if required: raise ValueError('Selecciona la columna: '+label)
        elif isinstance(column,bool) or not isinstance(column,int) or not 0<=column<len(table['headers']):
            raise ValueError('Columna no válida: '+label)
        else: columns.append(column)
    if len(set(columns))!=len(columns): raise ValueError('Una columna no puede asignarse a dos campos.')
    cutoff=source_date(cutoff)
    records=[];errors=[];seen=set();warnings=[]
    for row in table['rows']:
        output={};problems=[]
        for field,label,required in fields:
            column=mapping.get(field)
            if column is None and (field in {k for k,_ in EXTENSIONS.get(kind,[])} or field in STOCK_PARAMETER_KEYS or kind=='sales_line' and field in {k for k,_,_ in SALES_DETAILS} or kind=='cash_flow' and field in {k for k,_,_ in COLLECTION_LINKS}):
                continue  # Preserve older inventory payloads; absent parameters are not defaults.
            value=row['values'][column] if column is not None else None
            if value is None or (isinstance(value,str) and not value.strip()):
                if required: problems.append(label+': dato vacío')
                output[field]=None;continue
            try:
                if isinstance(value,dict):
                    if field not in CACHED_PRICE_FIELDS.get(kind,set()) or value.get('cached') is None:raise ValueError('La celda contiene una fórmula sin resultado utilizable para este campo. Recalcula y guarda el libro o aporta valores fijos.')
                    value=value['cached'];warnings.append('Fila '+str(row['row'])+': '+label+' utiliza el resultado guardado de una fórmula; no se ejecuta la fórmula ni se acredita su vigencia.')
                if field in NUMBERS: value=str(number(value,convention))
                elif field in ('issue_date','due_date','observation_date','valid_from','valid_until','reference_date','order_date','delivery_date','contractual_due_date'): value=source_date(value)
                elif field in IDENTIFIERS and not isinstance(value,str):
                    raise ValueError('El código debe ser texto para conservar sus ceros iniciales.')
                else: value=str(value).strip()
                output[field]=value
            except (ValueError,InvalidOperation) as error: problems.append(label+': '+str(error))
        if not problems:
            key_fields={'supplier':['supplier_id'],'customer':['customer_id'],'article':['article_id'],'inventory':['article_id','unit'],'stock_policy':['article_id','unit'],'cash_flow':['flow_id'],'balance':['account_code','currency'],'price_reference':['article_id','unit']}.get(kind)
            if key_fields:
                key=tuple(output.get(k) for k in key_fields)
                if key in seen: problems.append('Identificador duplicado dentro del archivo.')
                seen.add(key)
            if kind in ('purchase_line','sales_line'):
                fingerprint=tuple(sorted(output.items()))
                if fingerprint in seen: warnings.append('Fila '+str(row['row'])+': mismo contenido que otra línea de factura; verifica si es un duplicado. Se conservarán ambas filas.')
                seen.add(fingerprint)
                if output['issue_date']>cutoff: problems.append('La factura es posterior a la fecha de corte.')
                if Decimal(output['quantity'])<=0 or Decimal(output['net_amount'])<0: problems.append('Las líneas de compra o venta deben tener cantidad positiva e importe neto no negativo; separa los abonos.')
            if kind=='purchase_line' and output.get('unit_price') is not None and Decimal(output['quantity'])>0:
                if Decimal(output['unit_price'])<0 or abs(Decimal(output['unit_price'])-Decimal(output['net_amount'])/Decimal(output['quantity']))>Decimal('.005'):
                    problems.append('Precio unitario incompatible con el importe neto y la cantidad.')
            if kind=='price_reference':
                if Decimal(output['reference_cost'])<0 or Decimal(output['reference_price'])<0:problems.append('Las referencias de precio no pueden ser negativas.')
                if output.get('target_margin') is not None and not 0<=Decimal(output['target_margin'])<1:problems.append('El margen supuesto debe ser una fracción entre 0 y 1, sin incluir 1.')
            if kind=='sales_line':
                from eios.company_sales_receipt_math import validate_line
                problems.extend(validate_line(output,warnings,row['row']))
            if kind=='sales_installment':
                n=Decimal(output['installment_number'])
                if n<1 or n!=n.to_integral_value():problems.append('El número de cuota debe ser un entero positivo.')
                amount=Decimal(output['amount']);pending=output.get('pending_amount')
                if amount<0:problems.append('La cuota no puede ser negativa; separa los abonos.')
                if pending is not None and not 0<=Decimal(pending)<=amount:problems.append('El pendiente debe estar entre cero y el importe original de la cuota.')
                quota_key=(output['customer_id'],output['invoice_number'],str(n.normalize()))
                ref_key=('installment_ref',output['installment_ref'])
                if quota_key in seen or ref_key in seen:problems.append('Referencia de cuota o número de cuota repetido para la misma factura y cliente.')
                seen.update((quota_key,ref_key))
            if kind=='sale_price':
                if Decimal(output['recommended_sale_price'])<0:problems.append('El PVP no puede ser negativo.')
                if output['valid_until']<output['valid_from']:problems.append('La vigencia final no puede preceder a la inicial.')
            if kind=='cash_flow':
                output['direction']=output['direction'].upper()
                if output['direction'] not in ('COBRO','PAGO'): problems.append('El movimiento debe indicar COBRO o PAGO.')
                if Decimal(output['amount'])<0: problems.append('Usa un importe positivo e indica COBRO o PAGO.')
            if (kind!='price_reference' or output.get('currency') is not None) and 'currency' in output and not re.fullmatch('[A-Z]{3}',output['currency'] or ''): problems.append('Moneda: usa un código de tres letras mayúsculas, por ejemplo EUR.')
            if kind in ('inventory','stock_policy','balance') and output.get('observation_date') and output['observation_date']!=cutoff:
                problems.append('La fecha de referencia de la fila no coincide con la fecha de corte indicada.')
            if kind in ('inventory','stock_policy'):
                output['as_of_date']=cutoff
                supplied=[key for key in STOCK_PARAMETER_KEYS if output.get(key) is not None]
                if kind=='stock_policy' and not supplied:
                    problems.append('Aporta al menos un parámetro: mínimo, seguridad, punto de pedido o plazo de entrega.')
                for key in supplied:
                    if Decimal(output[key])<0: problems.append('No se admite un valor negativo en '+next(label for field,label,_ in STOCK_PARAMETERS if field==key)+'.')
                days=output.get('lead_time_days')
                if days is not None and Decimal(days)!=Decimal(days).to_integral_value():
                    problems.append('El plazo de entrega debe indicar días enteros.')
            if kind=='inventory':
                output['as_of_date']=cutoff
                output['available']=str(Decimal(output['physical'])-Decimal(output['reserved'])) if output['reserved'] is not None else None
                if output['reserved'] is None: warnings.append('Fila '+str(row['row'])+': sin reservas; no se calcula stock disponible.')
            if kind=='balance': output['as_of_date']=cutoff
        if problems: errors.append(dict(row=row['row'],messages=problems))
        else: records.append(dict(kind=kind,data=output,source_row=row['row']))
    return dict(records=records,errors=errors,error_count=len(errors),warnings=warnings)
