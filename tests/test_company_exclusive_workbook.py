from io import BytesIO
import sqlite3
import unittest
from openpyxl import Workbook
from tests.test_company_sales_negotiation import SalesNegotiationTest as Fixture
from eios.company_workbook_reception import CompanyWorkbookReception,HEADERS
from eios.company_exclusive_workbook import activate,options,dataset_options
from eios.company_reception import CompanyReception


def book(customer='4300001',extra=False):
    wb=Workbook();g=wb.active;g.title='97_Guia_Carga';g.append(HEADERS)
    tables={
        'Clientes':('customer',['customer_id','legal_name','tax_id'],[[customer,'Cliente del Excel','B12345678']]),
        'Proveedores':('supplier',['supplier_id','legal_name'],[['4000001','Proveedor del Excel']]),
        'Articulos':('article',['article_id','description','unit','supplier_id','family'],[['A','Artículo del Excel','palets','4000001','Sustratos']]),
        'Inventario':('inventory',['article_id','unit','physical','reserved','minimum_stock'],[['A','palets',10,2,3]]),
        'Ventas':('sales_line',['invoice_number','issue_date','customer_id','article_id','quantity','unit','net_amount','currency','vat_amount','invoice_total'],[['FV-NEW','2026-09-20',customer,'A',2,'palets',100,'EUR',21,121]]),
        'Cuotas':('sales_installment',['installment_ref','customer_id','invoice_number','installment_number','due_date','amount','currency','payment_method','pending_amount'],[['Q-NEW',customer,'FV-NEW',1,'2026-10-30',121,'EUR','Transferencia',121]]),
        'Balance':('balance',['account_code','balance','currency'],[['572',999,'EUR']]),
        'Tesoreria':('cash_flow',['flow_id','direction','amount','currency','due_date','customer_id','invoice_number','installment_ref'],[['C-NEW','COBRO',121,'EUR','2026-10-30',customer,'FV-NEW','Q-NEW']])}
    if extra:tables.pop('Ventas');tables.pop('Cuotas');tables.pop('Tesoreria')
    for sheet,(kind,headers,rows) in tables.items():
        sh=wb.create_sheet(sheet);sh.append(headers)
        for row in rows:sh.append(row)
        for field in headers:g.append([sheet,kind,field,field])
    out=BytesIO();wb.save(out);return out.getvalue()


class ExclusiveWorkbookTest(unittest.TestCase):
    def setUp(self):
        Fixture.setUp(self)
        with sqlite3.connect(self.path) as c:
            c.execute('DROP TABLE documentary_source')
            c.execute('CREATE TABLE documentary_source(dataset_id TEXT,document_ref TEXT,metadata_json TEXT,content BLOB,PRIMARY KEY(dataset_id,document_ref))')
    tearDown=Fixture.tearDown
    add=Fixture.add
    financial_data=Fixture.financial_data
    set_policy=Fixture.set_policy

    def receive(self,content=None):
        reader=CompanyWorkbookReception(self.path)
        p=reader.preview(content or book(),'Empresa.xlsx',dict(company_id='one',legal_name='Empresa sintética',as_of_date='2026-09-30',material_nature='SYNTHETIC',number_format='ES',preserve_revision=True))
        self.assertTrue(p['ticket'],p)
        reader.commit(p['ticket'])
        return options(self.path,'one')['files'][0]['batch_id']

    def test_exclusive_snapshot_never_inherits_customers_stock_history_or_balance(self):
        old=self.desk.state('one');bid=self.receive();r=activate(self.path,'one',bid);did=r['dataset_id']
        state=self.desk.sales_state(did)
        self.assertEqual([c['customer_id'] for c in state['customers']],['4300001'])
        self.assertEqual([s['supplier_id'] for s in state['suppliers']],['4000001'])
        self.assertEqual(state['articles'][0]['description'],'Artículo del Excel')
        self.assertEqual(state['inventory']['A']['available'],'8')
        self.assertEqual(state['stock_minimums']['A']['value'],'3')
        self.assertEqual(state['stock_minimums']['A']['source'],'WORKBOOK')
        import json
        json.dumps(state)
        self.assertEqual([f['invoice_number'] for f in state['sales']['invoices']],['FV-NEW'])
        self.assertEqual(state['purchases'],[]);self.assertEqual(state['sales_proposals'],[])
        self.assertEqual(state['sales_financial_context']['financial_snapshot']['available_treasury'],'999')
        self.assertEqual(state['sales_quality']['summary']['matched'],1)
        self.assertEqual(self.desk.state('one')['customers'],old['customers'])
        self.assertEqual(activate(self.path,did,bid)['dataset_id'],did)
        self.assertFalse(r['decision_authority'])
        with sqlite3.connect(self.path) as c:
            rows=c.execute('SELECT dataset_id,company_id,as_of_date FROM documentary_company_snapshot').fetchall()
        self.assertTrue(dataset_options(self.path,rows)[0]['active_workbook'])
        self.assertFalse(next(r for r in dataset_options(self.path,rows) if r['dataset_id']=='two')['historical'])

    def test_later_unrelated_financial_receipt_cannot_replace_pinned_workbook(self):
        bid=self.receive();did=activate(self.path,'one',bid)['dataset_id']
        reader=CompanyReception(self.path);file=reader.read(b'Cuenta;Saldo;Moneda\n572;999999;EUR\n','Otro.csv')
        p=reader.preview(dict(upload=file['upload'],company_id='one',legal_name='Empresa sintética',as_of_date='2026-09-30',material_nature='SYNTHETIC',number_format='ES',kind='balance',mapping=dict(account_code=0,balance=1,currency=2)))
        other=reader.commit(p['ticket'])['batch_id']
        finance=self.desk.sales_state(did)['sales_financial_context']
        self.assertEqual(finance['financial_snapshot']['available_treasury'],'999')
        from eios.company_sales_treasury import checked_source
        with sqlite3.connect(self.path) as c,self.assertRaisesRegex(ValueError,'Excel activo'):
            checked_source(c,did,other,'balance')

    def test_new_book_missing_history_does_not_restore_old_invoices(self):
        bid=self.receive(book(extra=True));did=activate(self.path,'one',bid)['dataset_id'];state=self.desk.sales_state(did)
        self.assertEqual(state['sales']['invoices'],[])
        self.assertEqual(state['history']['A'],[])
        self.assertIsNone(state['sales_financial_context']['working_capital_input'])
        self.assertEqual(state['sales_financial_context']['cash_flows'],[])
        self.assertEqual(state['info']['exclusive_workbook']['filename'],'Empresa.xlsx')

    def test_old_customer_is_rejected_and_corrupt_activation_is_atomic(self):
        bid=self.receive();did=activate(self.path,'one',bid)['dataset_id']
        state=self.desk.sales_state(did)
        args=dict(self.offer,dataset_id=did,context_fingerprint=state['documentary_context']['fingerprint'])
        with self.assertRaisesRegex(ValueError,'histórica'):
            self.desk.sales_offer(dict(self.offer,context_fingerprint=self.desk.state('one')['documentary_context']['fingerprint']),persist=False)
        with self.assertRaisesRegex(ValueError,'cliente'):
            self.desk.sales_offer(args,persist=False)
        with sqlite3.connect(self.path) as c:
            before=c.execute('SELECT count(*) FROM documentary_company_snapshot').fetchone()[0]
            c.execute('UPDATE received_company_batch SET content=? WHERE batch_id=?',(b'corrupt',bid))
        with self.assertRaisesRegex(ValueError,'integridad'):activate(self.path,'one',bid)
        with sqlite3.connect(self.path) as c:
            self.assertEqual(c.execute('SELECT count(*) FROM documentary_company_snapshot').fetchone()[0],before)
            self.assertEqual(c.execute('SELECT dataset_id FROM company_active_workbook WHERE company_id=?',('one',)).fetchone()[0],did)

    def test_price_reference_without_unit_does_not_invent_a_unit_cost(self):
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(book()))
        sheet=wb.create_sheet('Precios');headers=['article_id','reference_cost','reference_price','currency']
        sheet.append(headers);sheet.append(['A',10,20,'EUR'])
        for field in headers:wb['97_Guia_Carga'].append(['Precios','price_reference',field,field])
        output=BytesIO();wb.save(output)
        bid=self.receive(output.getvalue());did=activate(self.path,'one',bid)['dataset_id']
        state=self.desk.state(did)
        self.assertIsNone(state['articles'][0].get('catalogue_unit_cost'))
        self.assertIsNone(state['articles'][0].get('recommended_sale_price'))
        self.assertTrue(any(r['data']['table_kind']=='price_reference' for r in self.desk.store.records(did,'source_sheet_row')))
