from copy import deepcopy
from io import BytesIO
import unittest
from openpyxl import load_workbook
from eios.company_sales_tariff_template import build_template
from eios.company_tables import read_table,normalize


class TariffTemplateTest(unittest.TestCase):
    def state(self):
        return dict(info=dict(company=dict(company_id='DEMO-001',legal_name='Empresa de prueba'),as_of_date='2026-09-30'),documentary_context=dict(material_nature='SYNTHETIC'),articles=[dict(article_id='0007',description='Artículo de prueba',unit_as_documented='uds.'),dict(article_id='=SKU',description='=Texto de origen',unit_as_documented='palets')])

    def test_codes_stay_text_and_missing_prices_are_not_zero_or_formulas(self):
        state=self.state();original=deepcopy(state);r=build_template(state)
        workbook=load_workbook(BytesIO(r['content']));sheet=workbook['Tarifa']
        self.assertEqual(sheet['A2'].value,'0007');self.assertEqual(sheet['A2'].data_type,'s')
        self.assertEqual(sheet['A3'].value,'=SKU');self.assertEqual(sheet['A3'].data_type,'s');self.assertEqual(sheet['B3'].data_type,'s')
        self.assertIsNone(sheet['D2'].value);self.assertIsNone(sheet['F2'].value);self.assertIsNone(sheet['G2'].value)
        self.assertEqual(sheet['D2'].fill.fgColor.rgb,'00FFF2CC')
        self.assertEqual(workbook['Instrucciones']['B2'].value,'DEMO-001')
        self.assertEqual(state,original);self.assertEqual(r['article_count'],2)
        workbook.close()

    def test_completed_template_roundtrips_through_the_actual_tariff_reader(self):
        r=build_template(self.state());workbook=load_workbook(BytesIO(r['content']));sheet=workbook['Tarifa']
        sheet.delete_rows(3)
        sheet['D2']='12,50';sheet['F2']='01/10/2026';sheet['G2']='31/12/2026'
        out=BytesIO();workbook.save(out);workbook.close()
        table=read_table(out.getvalue(),r['filename'],sheet='Tarifa')
        mapping=dict(article_id=0,unit=2,recommended_sale_price=3,currency=4,valid_from=5,valid_until=6)
        result=normalize(table,'sale_price',mapping,'2026-09-30','ES')
        self.assertEqual(result['error_count'],0)
        self.assertEqual(result['records'][0]['data']['article_id'],'0007')
        self.assertEqual(result['records'][0]['data']['recommended_sale_price'],'12.50')
        blank=normalize(read_table(r['content'],r['filename'],sheet='Tarifa'),'sale_price',mapping,'2026-09-30','ES')
        self.assertEqual(blank['error_count'],2)

    def test_missing_units_or_ambiguous_skus_cannot_create_a_misleading_template(self):
        for articles in ([],[dict(article_id='A',unit_as_documented='')],[dict(article_id='A',unit_as_documented='uds.')]*2):
            with self.subTest(articles=articles),self.assertRaises(ValueError):build_template(dict(self.state(),articles=articles))
