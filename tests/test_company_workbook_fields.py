from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from tests.test_company_workbook_reception import workbook, OPTIONS, count
from eios.company_workbook_reception import CompanyWorkbookReception
from eios.company_workbook_fields import header_index
from eios.company_tables import normalize, read_table
from eios.company_received_articles import received_article


class WorkbookFieldsTest(unittest.TestCase):
    def test_price_reference_and_optional_fields_are_received_without_activation(self):
        wb=load_workbook(BytesIO(workbook()))
        wb['Suppliers'].append([])
        wb['Suppliers']['C1']='Correo electrónico';wb['Suppliers']['C2']='contact@example.com'
        wb['Articles']['E1']='Subfamilia';wb['Articles']['E2']='Accesorios'
        sheet=wb.create_sheet('08_Referencias_Precios')
        sheet.append(['Sku','Coste fuente','PVP fuente','Unidad','Origen del coste','Origen del PVP'])
        sheet.append(['A1',10,15,'Unidad','Supuesto ficticio','Supuesto ficticio'])
        out=BytesIO();wb.save(out)
        with tempfile.TemporaryDirectory() as tmp:
            db=Path(tmp)/'test.db';service=CompanyWorkbookReception(db)
            report=service.preview(out.getvalue(),'book.xlsx',OPTIONS)
            self.assertEqual(report['status'],'READY')
            saved=service.commit(report['ticket']);self.assertEqual(count(db),6)
            result=received_article(db,saved['batch_ids'],'A1')
            self.assertEqual(result['catalog'][0]['data']['subfamily'],'Accesorios')
            self.assertEqual(result['price_references'][0]['data']['reference_price'],'15')
            self.assertEqual(result['price_references'][0]['data']['cost_origin'],'Supuesto ficticio')
            self.assertFalse(result['decision_authority']);self.assertEqual(result['operational_path'],'FORBIDDEN')
            self.assertNotIn('recommended_sale_price',result['catalog'][0]['data'])

    def test_headers_accept_line_breaks_but_reject_ambiguity(self):
        self.assertEqual(header_index(['Número\n de línea'],'Número de línea'),0)
        with self.assertRaises(ValueError):header_index(['Sku','SKU'],'Sku')

    def test_only_scoped_cached_price_results_are_accepted(self):
        table=dict(headers=['Sku','Coste fuente','PVP fuente','Unidad'],rows=[dict(row=2,values=['A1',10,dict(formula=True,cached=15),'Unidad'])])
        mapping=dict(article_id=0,reference_cost=1,reference_price=2,unit=3)
        r=normalize(table,'price_reference',mapping,'2026-09-30','ES')
        self.assertEqual(r['error_count'],0);self.assertTrue(r['warnings'])
        table['rows'][0]['values'][2]=dict(formula=True,cached=None)
        self.assertEqual(normalize(table,'price_reference',mapping,'2026-09-30','ES')['error_count'],1)
        table['rows'][0]['values'][0]=dict(formula=True,cached='A1')
        self.assertEqual(normalize(table,'price_reference',mapping,'2026-09-30','ES')['error_count'],1)

    def test_purchase_price_must_agree_with_quantity_and_net(self):
        headers=['Factura','Fecha','Proveedor','Sku','Cantidad','Unidad','Neto','Moneda','Precio']
        mapping=dict(invoice_number=0,issue_date=1,supplier_id=2,article_id=3,quantity=4,unit=5,net_amount=6,currency=7,unit_price=8)
        values=['I1','01/09/2026','S1','A1',2,'Unidad',20,'EUR',dict(formula=True,cached=10)]
        table=dict(headers=headers,rows=[dict(row=2,values=values)])
        self.assertEqual(normalize(table,'purchase_line',mapping,'2026-09-30','ES')['error_count'],0)
        values[-1]=11
        self.assertEqual(normalize(table,'purchase_line',mapping,'2026-09-30','ES')['error_count'],1)
        values[4]=0
        self.assertEqual(normalize(table,'purchase_line',mapping,'2026-09-30','ES')['error_count'],1)

    def test_revision_preserves_previous_receipts_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as tmp:
            db=Path(tmp)/'test.db';service=CompanyWorkbookReception(db)
            first=service.preview(workbook(),'book.xlsx',OPTIONS);service.commit(first['ticket'])
            changed=workbook(description='Changed article')
            self.assertEqual(service.preview(changed,'book.xlsx',OPTIONS)['status'],'CONFLICT')
            report=service.preview(changed,'book.xlsx',dict(OPTIONS,preserve_revision=True))
            self.assertEqual(report['status'],'READY');service.commit(report['ticket'])
            self.assertEqual(count(db),6)
            again=service.preview(changed,'book.xlsx',dict(OPTIONS,preserve_revision=True))
            self.assertEqual(service.commit(again['ticket'])['batch_ids'],[])
            self.assertEqual(count(db),6)

    def test_formula_without_cache_is_visible_to_the_reader(self):
        wb=load_workbook(BytesIO(workbook()));wb['Articles']['B2']='="Description"';out=BytesIO();wb.save(out)
        table=read_table(out.getvalue(),'book.xlsx',sheet='Articles')
        self.assertTrue(table['rows'][0]['values'][1]['formula'])
        self.assertIsNone(table['rows'][0]['values'][1]['cached'])
