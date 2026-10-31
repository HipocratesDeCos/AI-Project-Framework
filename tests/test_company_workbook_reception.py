from io import BytesIO
import sqlite3
import unittest
import tempfile
from pathlib import Path
from openpyxl import Workbook

from eios.company_workbook_reception import CompanyWorkbookReception, HEADERS

OPTIONS = dict(company_id='SAMPLE', legal_name='Sample Ltd', material_nature='SYNTHETIC',
               as_of_date='2026-09-30', number_format='ES')


def workbook(description='Test article', missing=False, invalid=False, ambiguous=False):
    wb = Workbook(); guide = wb.active; guide.title='97_Guia_Carga'; guide.append(HEADERS)
    tables={'Suppliers': ('supplier', ['supplier_id','legal_name'], [['S1','Supplier']]),
            'Articles': ('article', ['article_id','description','unit','supplier_id'], [['A1',description,'Unidad','absent' if missing else 'S1']]),
            'Stock': ('inventory', ['article_id','unit','physical','reserved'], [['A1','Unidad','invalid' if invalid else 10,None]])}
    for name,(kind,headers,rows) in tables.items():
        sheet=wb.create_sheet(name);sheet.append(headers)
        for row in rows:sheet.append(row)
        for header in headers:guide.append([name,kind,header,header])
    if ambiguous:wb['Articles'].cell(1,5,'article_id')
    wb.create_sheet('Notes').append(['Preserve this original support'])
    output=BytesIO();wb.save(output);return output.getvalue()


def count(db):
    with sqlite3.connect(db) as c:
        if not c.execute("SELECT 1 FROM sqlite_master WHERE name='received_company_batch'").fetchone():return 0
        return c.execute('SELECT count(*) FROM received_company_batch').fetchone()[0]


class WorkbookReceptionTest(unittest.TestCase):
    def setUp(self):
        temp=tempfile.TemporaryDirectory();self.addCleanup(temp.cleanup);self.tmp_path=Path(temp.name)

    def test_preview_closes_all_database_connections(self):
        from unittest.mock import patch
        connections = []
        class TrackedConnection(sqlite3.Connection):
            was_closed = False
            def close(self):
                self.was_closed = True
                return super().close()
        original = sqlite3.connect
        def connect(*args, **kwargs):
            kwargs['factory'] = TrackedConnection
            connection = original(*args, **kwargs)
            connections.append(connection)
            return connection
        with patch('sqlite3.connect', side_effect=connect):
            report = CompanyWorkbookReception(self.tmp_path/'close.db').preview(workbook(), 'book.xlsx', OPTIONS)
        self.assertTrue(report['ticket'])
        self.assertTrue(connections)
        self.assertTrue(all(c.was_closed for c in connections))

    def test_batch_idempotent_and_original_preserved(self):
        tmp_path=self.tmp_path
        db=tmp_path/'company.db';service=CompanyWorkbookReception(db);content=workbook()
        report=service.preview(content,'book.xlsx',OPTIONS)
        assert count(db)==0 and report['ticket'] and 'Notes' in report['ignored_sheets']
        assert report['financial_reconciliation']=='NOT_PERFORMED'
        saved=service.commit(report['ticket']);assert len(saved['batch_ids'])==3 and count(db)==3
        with sqlite3.connect(db) as c:
            assert all(bytes(row[0])==content for row in c.execute('SELECT content FROM received_company_batch'))
            assert not c.execute("SELECT 1 FROM sqlite_master WHERE name='documentary_company_record'").fetchone()
        report=service.preview(content,'book.xlsx',OPTIONS)
        assert service.commit(report['ticket'])['status']=='UNCHANGED' and count(db)==3
        report=service.preview(content,'renamed.xlsx',OPTIONS)
        assert all(row['status']=='UNCHANGED' and row['new_source_version'] for row in report['comparison'])
        assert service.commit(report['ticket'])['status']=='RECEIVED' and count(db)==6


    def test_bad_sheet_or_reference_never_writes(self):
        tmp_path=self.tmp_path
        kwargs={'missing':True}
        db=tmp_path/'company.db';service=CompanyWorkbookReception(db)
        assert service.preview(workbook(**kwargs),'book.xlsx',OPTIONS)['ticket'] is None
        assert count(db)==0

    def test_invalid_values_and_corrupt_file(self):
        service=CompanyWorkbookReception(self.tmp_path/'company.db')
        report=service.preview(workbook(invalid=True),'book.xlsx',OPTIONS)
        assert report['status']=='INVALID' and not report['ticket']
        with self.assertRaisesRegex(ValueError,'guía'):
            service.preview(b'not a zip','book.xlsx',OPTIONS)
        assert count(self.tmp_path/'company.db')==0


    def test_conflict_does_not_overwrite(self):
        tmp_path=self.tmp_path
        db=tmp_path/'company.db';service=CompanyWorkbookReception(db)
        service.commit(service.preview(workbook(),'book.xlsx',OPTIONS)['ticket'])
        report=service.preview(workbook(description='Different article'),'new.xlsx',OPTIONS)
        assert report['status']=='CONFLICT' and not report['ticket'] and count(db)==3


    def test_stale_and_expired_reviews(self):
        tmp_path=self.tmp_path
        db=tmp_path/'company.db';now=[0];a=CompanyWorkbookReception(db,clock=lambda:now[0]);b=CompanyWorkbookReception(db)
        ticket=a.preview(workbook(),'a.xlsx',OPTIONS)['ticket']
        b.commit(b.preview(workbook(),'b.xlsx',OPTIONS)['ticket'])
        with self.assertRaisesRegex(ValueError,'ha cambiado'):a.commit(ticket)
        ticket=a.preview(workbook(),'a.xlsx',OPTIONS)['ticket'];now[0]=601
        with self.assertRaisesRegex(ValueError,'caducado'):a.commit(ticket)
        assert count(db)==3


    def test_ambiguous_column_and_company_identity(self):
        tmp_path=self.tmp_path
        db=tmp_path/'company.db';service=CompanyWorkbookReception(db)
        with self.assertRaisesRegex(ValueError,'ambigua'):service.preview(workbook(ambiguous=True),'book.xlsx',OPTIONS)
        service.commit(service.preview(workbook(),'book.xlsx',OPTIONS)['ticket'])
        with self.assertRaisesRegex(ValueError,'razón social'):service.preview(workbook(),'book.xlsx',dict(OPTIONS,legal_name='Other'))


    def test_transaction_rolls_back_all_sheets(self):
        tmp_path=self.tmp_path
        db=tmp_path/'company.db';service=CompanyWorkbookReception(db)
        ticket=service.preview(workbook(),'book.xlsx',OPTIONS)['ticket']
        # Force a failure after the first insert without modifying production code.
        with sqlite3.connect(db) as c:
            c.executescript('''CREATE TABLE received_company(company_id TEXT PRIMARY KEY,legal_name TEXT,material_nature TEXT);
              CREATE TABLE received_company_batch(batch_id TEXT PRIMARY KEY,company_id TEXT,received_at TEXT,metadata_json TEXT,content BLOB,records_json TEXT);
              CREATE TRIGGER fail_second BEFORE INSERT ON received_company_batch WHEN (SELECT count(*) FROM received_company_batch)>0
              BEGIN SELECT RAISE(ABORT,'injected write failure'); END;''')
        with self.assertRaises(sqlite3.IntegrityError):service.commit(ticket)
        assert count(db)==0
